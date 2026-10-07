"""Создаёт шаблон прайс-листа для импорта в Google Таблицы."""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

root = Path(sys.argv[1])
rows = json.loads((root / "tools" / "prices.json").read_text(encoding="utf-8"))
out_path = root / "price-template.xlsx"

ACCENT = "147746"
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
CAT_FILL = PatternFill("solid", fgColor="E8F3ED")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="BFC9C3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()

# ------------------------------------------------------------------ лист «Цены»
ws = wb.active
ws.title = "Цены"

header = rows[0]
for col, title in enumerate(header, start=1):
    cell = ws.cell(row=1, column=col, value=title)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER

for r, row in enumerate(rows[1:], start=2):
    for c, value in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=value)
        cell.border = BORDER
        cell.alignment = Alignment(
            horizontal="left" if c <= 3 else "center", vertical="top", wrap_text=(c == 3)
        )
        if c == 1:
            cell.font = Font(bold=True, color="0F5132")
            cell.fill = CAT_FILL
        if c in (4, 5, 6, 7, 8, 9):
            cell.font = Font(bold=True)

ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows)}"

widths = {1: 30, 2: 46, 3: 44, 4: 12, 5: 12, 6: 14, 7: 13, 8: 15, 9: 14, 10: 44}
for col, width in widths.items():
    ws.column_dimensions[get_column_letter(col)].width = width
ws.row_dimensions[1].height = 34

# ------------------------------------------------------- лист «Как заполнять»
help_ws = wb.create_sheet("Как заполнять")
help_ws.column_dimensions["A"].width = 118
help_ws["A1"] = "Как вести прайс-лист на сайте CGP-DETAILING"
help_ws["A1"].font = Font(bold=True, size=15, color=ACCENT)
help_ws["A1"].alignment = Alignment(vertical="center")
help_ws.row_dimensions[1].height = 26

lines = [
    "",
    "Лист «Цены» — это и есть прайс-лист на сайте. Что здесь написано, то и видят клиенты.",
    "Правки в таблице попадают на сайт автоматически, обычно в течение 5 минут.",
    "",
    "КОЛОНКИ",
    "• Категория — раздел прайса (кнопка-вкладка на сайте). Одинаковое название = один раздел.",
    "  Порядок разделов на сайте = порядок, в котором категории впервые встречаются сверху вниз.",
    "• Описание раздела — подзаголовок под названием раздела. Достаточно заполнить в одной",
    "  строке раздела, сайт возьмёт первое непустое значение.",
    "• Услуга — название услуги в таблице. Обязательная колонка.",
    "• Хэтчбек / Седан / Бизнес-седан / Кроссовер / Минивэн / Джип — цены по типам авто.",
    "  Если у услуги цена не зависит от типа авто — заполните только колонку «Цена»,",
    "  тогда услуга попадёт в блок «Дополнительно» с одной ценой.",
    "• Цена — цена для услуг, которые не делятся по типам авто.",
    "• Примечание — мелкая серая подпись под названием услуги. Можно оставить пустой.",
    "",
    "КАК ДОБАВИТЬ НОВУЮ УСЛУГУ",
    "Просто добавьте строку внизу раздела: повторите категорию, впишите название и цену.",
    "Ничего больше делать не нужно — услуга появится на сайте сама.",
    "",
    "КАК ИЗМЕНИТЬ ЦЕНУ",
    "Щёлкните ячейку и впишите новое значение. Пишите либо просто число (500),",
    "либо текст (от 500, 800–900, по запросу) — на сайте отобразится ровно это.",
    "Просто число сайт покажет как «500 р.».",
    "",
    "КАК ДОБАВИТЬ НОВЫЙ ТИП АВТОМОБИЛЯ",
    "Впишите название в свободную колонку справа (например, «Пикап») и заполните цены.",
    "Колонка появится в таблицах на сайте автоматически.",
    "",
    "КАК СДЕЛАТЬ НОВЫЙ РАЗДЕЛ",
    "Напишите новое название в колонке «Категория» у нужных строк. Появится новая вкладка.",
    "",
    "УДАЛИТЬ УСЛУГУ ИЛИ РАЗДЕЛ",
    "Удалите строки целиком. Пустые строки на сайте не показываются.",
    "",
    "ВАЖНО, ЧТОБЫ ЦЕНЫ ОБНОВЛЯЛИСЬ НА САЙТЕ",
    "Таблица должна быть опубликована в интернете:",
    "Файл → Поделиться → Опубликовать в интернете →",
    "слева выбрать лист «Цены», справа формат «CSV» → Опубликовать → скопировать ссылку.",
    "Эту ссылку один раз вставляют в файл index.html в настройку PRICE_CONFIG.csvUrl.",
    "После этого доступ к таблице можно не открывать — публикация уже работает.",
    "",
    "Если случайно испортить таблицу, сайт покажет встроенный прайс-лист из index.html,",
    "то есть без цен сайт не останется.",
]
for i, line in enumerate(lines, start=2):
    cell = help_ws.cell(row=i, column=1, value=line)
    if line and line == line.upper() and not line.startswith("•"):
        cell.font = Font(bold=True, color="0F5132")
    else:
        cell.font = Font(size=11)
    cell.alignment = Alignment(wrap_text=False, vertical="center")

wb.save(out_path)
print("saved:", out_path)
print("rows in Цены:", len(rows))
